"""
Lector de CERTIFICACIONES a cliente (formato SIS/Crystal Reports: ORIGEN · ANTERIOR · ACTUAL).

No usa IA: el informe tiene columnas fijas y marcadores 'OR', 'AN', 'AC', así que se
lee de forma determinista y se VERIFICA con su propia aritmética:
  1) en cada partida:     importe origen  = importe anterior + importe actual
  2) en cada partida:     importe origen ≈ cantidad origen × precio   (redondeo a 3 decimales)
  3) en cada (sub)capítulo: Σ partidas = total impreso del capítulo (origen, anterior, actual)
  4) en el documento:     Σ capítulos = 'Totales' impresos
Si algo no cuadra, se informa: nunca se da por buena una lectura que no cuadra.

Los importes de la certificación vienen con 3 decimales: se guardan en MILÉSIMAS de euro
(enteros) para no perder precisión.
"""
from __future__ import annotations

import io
import re
from dataclasses import dataclass, field
from datetime import datetime
from decimal import Decimal

import pdfplumber

from .money import parse_amount

NUM = r"-?(?:\d{1,3}(?:\.\d{3})+|\d+),\d{3}"
UNIDADES = {"m", "m2", "m3", "ml", "m²", "m³", "ud", "u", "uds", "un", "kg", "t", "tn", "h", "pa", "p.a.", "l", "mes",
            "d", "km", "jor", "par", "pa.", "ud.", "m.", "m2.", "m3.", "h.", "kg.", "%"}

RE_LINEA = re.compile(
    rf"^(?P<codigo>\S+)\s+(?P<resto>.*?)\s*(?P<pct>{NUM})%OR\s+(?P<cant>{NUM})(?:\s+(?P<precio>{NUM})\s+(?P<imp>{NUM}))?"
    rf"\s*AN\s*(?P<an>.*?)\s*AC\s*(?P<ac>.*?)\s*$")
RE_TOTAL = re.compile(rf"^\s*Total (?P<tipo>Capítulo|SubCapítulo) (?P<codigo>\S+) - (?P<nombre>.*?)\s*(?P<nums>(?:{NUM}\s*)*)$")
RE_TOTALES = re.compile(rf"^\s*Totales\s+(?P<nums>(?:{NUM}\s*€?\s*)+)$")
RE_CABECERA = re.compile(r"^(?P<codigo>[0-9A-Z][0-9A-Za-z._\-]*)\s{2,}(?P<nombre>[^\d\s].*?)\s*$")
RUIDO = ("LLORCA GROUP HISPANIA", "CERTIFICACIÓN", "Presupuesto:", "Certificación nº", "ORIGEN", "Código",
         "Página ", "En BENIDORM", "Propiedad", "Dirección Facultativa")


def M(v) -> int:
    """Decimal -> milésimas (entero)."""
    return int((Decimal(v) * 1000).to_integral_value())


def mil_a_dec(m: int | None) -> Decimal:
    return Decimal(int(m or 0)) / 1000


def _nums(s: str) -> list[Decimal]:
    return [parse_amount(x) for x in re.findall(NUM, s or "")]


def _limpia(s: str) -> str:
    """Quita el espaciado de 'letras sueltas' que genera Crystal: 'D E S M O N T E ,X' -> 'DESMONTE ,X'."""
    s = re.sub(r"\b(?:\w ){2,}\w\b", lambda m: m.group(0).replace(" ", ""), s)
    return re.sub(r"\s{2,}", " ", s).strip()


@dataclass
class Linea:
    codigo: str
    unidad: str | None
    titulo: str
    capitulo: str
    subcapitulo: str | None
    pct_origen: Decimal
    cant_origen: Decimal
    precio: Decimal
    imp_origen: Decimal
    cant_anterior: Decimal
    imp_anterior: Decimal
    cant_actual: Decimal
    imp_actual: Decimal
    descripcion: str = ""
    orden: int = 0

    @property
    def cant_presupuesto(self) -> Decimal | None:
        """Cantidad total prevista deducida de '% a origen' (cant origen / %)."""
        if self.pct_origen and self.pct_origen != 0:
            return (self.cant_origen / self.pct_origen * 100).quantize(Decimal("0.001"))
        return None

    @property
    def imp_presupuesto(self) -> Decimal | None:
        c = self.cant_presupuesto
        return None if c is None else (c * self.precio).quantize(Decimal("0.001"))


@dataclass
class Capitulo:
    codigo: str
    nombre: str
    padre: str | None
    nivel: int
    orden: int
    tot_origen: Decimal | None = None
    tot_anterior: Decimal | None = None
    tot_actual: Decimal | None = None


@dataclass
class Certificacion:
    numero: int | None = None
    fecha: str | None = None
    presupuesto: str | None = None
    obra: str | None = None
    cliente: str | None = None
    capitulos: list[Capitulo] = field(default_factory=list)
    lineas: list[Linea] = field(default_factory=list)
    totales: tuple | None = None
    avisos: list[str] = field(default_factory=list)
    paginas: int = 0


def clasificar_capitulo(codigo: str, nombre: str) -> str:
    n = (nombre or "").upper()
    if "ORDEN" in n and "CAMBIO" in n:
        return "orden_cambio"
    if "OPCIONAL" in n:
        return "opcional"
    if "REVISI" in n or "MODIFICAC" in n:
        return "modificacion"
    return "contrato"


def leer_texto(data: bytes) -> tuple[str, int]:
    """Texto con disposición (layout) — imprescindible para respetar el orden de columnas."""
    with pdfplumber.open(io.BytesIO(data)) as pdf:
        partes = [p.extract_text(layout=True, x_density=4.2) or "" for p in pdf.pages]
        return "\n".join(partes), len(pdf.pages)


def parsear(texto: str, paginas: int = 0) -> Certificacion:
    c = Certificacion(paginas=paginas)
    m = re.search(r"Presupuesto:\s*(\S+)\s+Obra:\s*(.*?)\s{2,}Fecha:\s*(\d{2}/\d{2}/\d{4})", texto)
    if m:
        c.presupuesto, c.obra = m.group(1), m.group(2).strip()
        c.fecha = datetime.strptime(m.group(3), "%d/%m/%Y").date().isoformat()
    m = re.search(r"Certificación nº:\s*(\d+)\s+Cliente:\s*(.*?)\s*$", texto, re.M)
    if m:
        c.numero, c.cliente = int(m.group(1)), m.group(2).strip()

    capitulos_cod = set(re.findall(r"Total Capítulo (\S+) - ", texto))
    subcap_cod = set(re.findall(r"Total SubCapítulo (\S+) - ", texto))
    caps: dict[str, Capitulo] = {}
    cap_actual, pila = None, []          # pila de subcapítulos abiertos (pueden anidarse)
    ultima: Linea | None = None
    esperando_titulo = False
    orden = 0

    for raw in texto.splitlines():
        s = raw.rstrip()
        st = s.strip()
        if not st or any(st.startswith(r) for r in RUIDO) or re.match(r"^Página \d+ de \d+$", st):
            continue

        mt = RE_TOTALES.match(s)
        if mt:
            c.totales = tuple(_nums(mt.group("nums")))
            continue
        mt = RE_TOTAL.match(s)
        if mt:
            cod = mt.group("codigo")
            n = _nums(mt.group("nums"))
            if cod in caps:
                caps[cod]._impresos = n  # se resuelven al final (columnas vacías se omiten en el PDF)
            if mt.group("tipo") == "SubCapítulo":
                if cod in pila:
                    del pila[pila.index(cod):]
            else:
                cap_actual, pila = None, []
            ultima, esperando_titulo = None, False
            continue

        ml = RE_LINEA.match(st)
        if ml:
            resto = ml.group("resto").strip()
            partes = resto.split(None, 1)
            unidad = None
            if partes and partes[0].lower() in UNIDADES and len(partes) > 1:
                unidad, resto = partes[0], partes[1]
            an, ac = _nums(ml.group("an")), _nums(ml.group("ac"))
            sin_precio = ml.group("precio") is None
            precio_txt, cant_txt = ml.group("precio"), ml.group("cant")
            if sin_precio and parse_amount(ml.group("pct")) == 0 and len(an) + len(ac) >= 2:
                # partida anulada: origen 0 (no se imprime), el único número es el precio
                precio_txt, cant_txt, sin_precio = cant_txt, "0", False
            if sin_precio:   # partida de medición sin precio (incluida en otra): solo cantidades
                an, ac = an + [Decimal(0)] if an else an, ac + [Decimal(0)] if ac else ac
            if len(an) not in (0, 2) or len(ac) not in (0, 2):
                c.avisos.append(f"Partida {ml.group('codigo')}: columnas anterior/actual irregulares")
            orden += 1
            ultima = Linea(
                codigo=ml.group("codigo"), unidad=unidad, titulo=_limpia(resto),
                capitulo=cap_actual or "?", subcapitulo=pila[-1] if pila else None,
                pct_origen=parse_amount(ml.group("pct")), cant_origen=parse_amount(cant_txt),
                precio=parse_amount(precio_txt or "0"), imp_origen=parse_amount(ml.group("imp") or "0"),
                cant_anterior=an[0] if an else Decimal(0), imp_anterior=an[1] if len(an) > 1 else Decimal(0),
                cant_actual=ac[0] if ac else Decimal(0), imp_actual=ac[1] if len(ac) > 1 else Decimal(0),
                orden=orden)
            c.lineas.append(ultima)
            esperando_titulo = True
            continue

        mc = RE_CABECERA.match(st)
        if mc and not re.search(NUM, st) and mc.group("codigo") in (capitulos_cod | subcap_cod):
            cod, nombre = mc.group("codigo"), _limpia(mc.group("nombre"))
            orden += 1
            if cod in capitulos_cod:
                cap_actual, pila = cod, []
                caps[cod] = Capitulo(cod, nombre, None, 1, orden)
            else:
                caps[cod] = Capitulo(cod, nombre, pila[-1] if pila else cap_actual, 2 + len(pila), orden)
                pila.append(cod)
            ultima, esperando_titulo = None, False
            continue

        if ultima is not None:
            # continuación: 1ª línea corta en mayúsculas = resto del título; lo demás = descripción larga
            if esperando_titulo and len(st) <= 45 and st.upper() == st:
                ultima.titulo = _limpia(ultima.titulo + " " + st)
            else:
                ultima.descripcion = (ultima.descripcion + " " + st).strip()[:4000]
            esperando_titulo = False

    # -------- resolver totales impresos por capítulo (columnas vacías no se imprimen)
    for cod, cap in caps.items():
        if cap.nivel == 1:
            ls = [l for l in c.lineas if l.capitulo == cod]
        else:
            ls = [l for l in c.lineas if _desciende(l.subcapitulo, cod, caps)]
        so = sum((l.imp_origen for l in ls), Decimal(0))
        sa = sum((l.imp_anterior for l in ls), Decimal(0))
        sc = sum((l.imp_actual for l in ls), Decimal(0))
        cap.tot_origen, cap.tot_anterior, cap.tot_actual = so, sa, sc
        imp = getattr(cap, "_impresos", None)
        if imp is None:
            c.avisos.append(f"Capítulo {cod}: sin total impreso")
            continue
        esperado = [x for x in (so, sa, sc) if x != 0]
        if [round(x, 2) for x in imp if x != 0] != [round(x, 2) for x in esperado]:
            c.avisos.append(f"Capítulo {cod} «{cap.nombre}»: Σ partidas {esperado} ≠ total impreso {imp}")
        del cap._impresos
    c.capitulos = sorted(caps.values(), key=lambda x: x.orden)
    return c


def _desciende(sub: str | None, cod: str, caps: dict) -> bool:
    while sub:
        if sub == cod:
            return True
        k = caps.get(sub)
        sub = k.padre if k and k.nivel > 1 else None
    return False


def verificar(c: Certificacion) -> list[tuple[str, str]]:
    """Devuelve [(severidad, mensaje)]. Lista vacía = lectura perfecta."""
    out = [("alta", a) for a in c.avisos]
    tol = Decimal("0.002")
    for l in c.lineas:
        if abs(l.imp_anterior + l.imp_actual - l.imp_origen) > tol:
            out.append(("alta", f"{l.codigo}: origen {l.imp_origen} ≠ anterior {l.imp_anterior} + actual {l.imp_actual}"))
        if abs(l.cant_anterior + l.cant_actual - l.cant_origen) > tol:
            out.append(("media", f"{l.codigo}: cantidades origen ≠ anterior + actual"))
        teor = (l.cant_origen * l.precio).quantize(Decimal("0.001"))
        if abs(teor - l.imp_origen) > Decimal("0.01") + abs(l.cant_origen) * Decimal("0.0005"):
            out.append(("media", f"{l.codigo}: cantidad × precio = {teor} ≠ importe {l.imp_origen}"))
        if l.capitulo == "?":
            out.append(("alta", f"{l.codigo}: partida fuera de capítulo"))
    if c.totales:
        so = sum((l.imp_origen for l in c.lineas), Decimal(0))
        sa = sum((l.imp_anterior for l in c.lineas), Decimal(0))
        sc = sum((l.imp_actual for l in c.lineas), Decimal(0))
        esperado = [so, sa, sc]
        if [round(x, 2) for x in c.totales] != [round(x, 2) for x in esperado][:len(c.totales)]:
            out.append(("critica", f"Totales impresos {list(c.totales)} ≠ Σ partidas {esperado}"))
    else:
        out.append(("alta", "No se encontró la línea de 'Totales' del documento"))
    if not c.numero:
        out.append(("alta", "No se ha leído el número de certificación"))
    return out


def leer_pdf(data: bytes) -> Certificacion:
    texto, n = leer_texto(data)
    return parsear(texto, n)


# ============================================================================ certificación en Excel / CSV
CAMPOS_TABLA = {
    "codigo": r"^c[oó]d", "unidad": r"^(ud|uds|unidad|u\.?)$", "descripcion": r"descrip|concepto|resumen|partida",
    "pct_origen": r"%|porcent", "cant_origen": r"cant.*(origen|or)|^cant|^medici", "precio": r"precio|p\.?\s?unit",
    "imp_origen": r"(imp|importe).*(origen|or\b)|^imp|^importe|^total", "imp_anterior": r"anterior|\ban\b",
    "imp_actual": r"actual|\bac\b|mes|periodo",
}


def _texto_celdas(filas: list[list]) -> str:
    return "\n".join(" ".join(str(x) for x in f if x not in (None, "")) for f in filas[:30])


def _num(v) -> Decimal | None:
    if v is None or v == "":
        return None
    if isinstance(v, (int, float, Decimal)):
        return Decimal(str(v))
    try:
        return parse_amount(str(v)) if re.search(r"\d", str(v)) else None
    except Exception:
        return None


def detectar_tabla(filas: list[list]) -> dict:
    """Devuelve {'modo': 'sis'|'cabecera', 'cols': {campo: índice}, 'inicio': fila}."""
    # 1) formato SIS/Crystal exportado a Excel: columna de marcas «OR»/«AN»/«AC»
    for j in range(min(40, max((len(f) for f in filas), default=0))):
        n_or = sum(1 for f in filas if j < len(f) and str(f[j]).strip() == "OR")
        if n_or >= 5:
            cols = {"marca": j}
            # código: columna con más valores tipo 01.02.03; descripción: la de textos largos en esas filas
            filas_or = [f for f in filas if j < len(f) and str(f[j]).strip() == "OR"]
            def mejor(pred):
                cuenta = {}
                for f in filas_or:
                    for k, v in enumerate(f[:j]):
                        if v not in (None, "") and pred(v):
                            cuenta[k] = cuenta.get(k, 0) + 1
                return max(cuenta, key=cuenta.get) if cuenta else None
            cols["codigo"] = mejor(lambda v: isinstance(v, str) and bool(re.fullmatch(r"[\w.\- /]{1,20}", v.strip()))
                                   and bool(re.search(r"\d|[A-Z]", v)) and v.strip().upper() not in ("OR", "AN", "AC"))
            cols["descripcion"] = mejor(lambda v: isinstance(v, str) and len(v) > 12)
            cols["unidad"] = mejor(lambda v: isinstance(v, str) and str(v).strip().lower() in UNIDADES)
            cols["pct_origen"] = mejor(lambda v: isinstance(v, (int, float)) and 0 <= v <= 100)
            if cols.get("pct_origen") == cols.get("codigo"):
                cols["pct_origen"] = None
            nums = [k for k in range(j + 1, j + 6) if sum(1 for f in filas_or if k < len(f) and isinstance(f[k], (int, float))) >= 3]
            if len(nums) >= 3:
                cols["cant_origen"], cols["precio"], cols["imp_origen"] = nums[0], nums[1], nums[-1]
            return {"modo": "sis", "cols": cols, "inicio": 0}
    # 2) tabla con cabecera (Presto, SIS «listado», hojas propias)
    for i, f in enumerate(filas[:40]):
        cab = [str(x or "").strip().lower() for x in f]
        cols = {}
        for campo, pat in CAMPOS_TABLA.items():
            for k, h in enumerate(cab):
                if h and re.search(pat, h) and k not in cols.values():
                    cols[campo] = k
                    break
        if "codigo" in cols and "descripcion" in cols and ("imp_origen" in cols or "cant_origen" in cols):
            return {"modo": "cabecera", "cols": cols, "inicio": i + 1}
    return {"modo": None, "cols": {}, "inicio": 0}


def leer_tabla(filas: list[list], mapeo: dict | None = None) -> Certificacion:
    """Certificación a partir de una hoja: líneas con código+precio = partidas; código sin importes = (sub)capítulo."""
    c = Certificacion()
    txt = _texto_celdas(filas)
    m = re.search(r"Certificaci[oó]n\s*n[º°o]?\s*[:.]?\s*(\d+)", txt, re.I)
    if m:
        c.numero = int(m.group(1))
    m = re.search(r"Presupuesto:?\s*(\d{4,})", txt)
    if m:
        c.presupuesto = m.group(1)
    m = re.search(r"Obra:?\s*([^\n]{3,60}?)\s{2,}|Obra:?\s*([^\n]{3,60})", txt)
    if m:
        c.obra = (m.group(1) or m.group(2) or "").strip()
    m = re.search(r"Cliente:?\s*([^\n]{3,80})", txt)
    if m:
        c.cliente = re.split(r"\s{2,}|(?<=S\.L\.)|(?<=S\.A\.)", m.group(1).strip())[0].strip()
    m = re.search(r"Fecha:?\s*(\d{1,2})[/.-](\d{1,2})[/.-](\d{4})", txt)
    if m:
        c.fecha = f"{m.group(3)}-{int(m.group(2)):02d}-{int(m.group(1)):02d}"
    for f in filas[:30]:
        for x in f:
            if isinstance(x, datetime):
                c.fecha = x.date().isoformat()
    det = mapeo or detectar_tabla(filas)
    cols = det["cols"]
    if not det.get("modo") or "codigo" not in cols:
        c.avisos.append("No se ha reconocido la estructura de la hoja: indique las columnas manualmente.")
        return c
    g = lambda f, k: f[cols[k]] if k in cols and cols[k] is not None and cols[k] < len(f) else None  # noqa: E731
    # capítulos: los que tienen fila «Total Capítulo X» (como en el PDF); si la hoja no las trae, por forma del código
    cods_cap = set()
    for f in filas:
        for x in f:
            if isinstance(x, str):
                mm = re.search(r"Total\s+Cap[ií]tulo\s+(\S+)", x)
                if mm:
                    cods_cap.add(mm.group(1))
    cap, sub, orden = None, None, 0
    caps = {}
    for i, f in enumerate(filas[det["inicio"]:], start=det["inicio"]):
        cod = str(g(f, "codigo") or "").strip()
        desc = str(g(f, "descripcion") or "").strip()
        if not cod or re.match(r"(?i)total|c[oó]digo", cod):
            continue
        es_partida = (str(g(f, "marca") or "").strip() == "OR") if det["modo"] == "sis" else \
            (_num(g(f, "imp_origen")) is not None or _num(g(f, "precio")) is not None)
        if not es_partida:
            if desc and not re.search(r"\d+[,.]\d{2}\s*$", desc):
                orden += 1
                if (cod in cods_cap) if cods_cap else ("." not in cod and "-" not in cod and len(cod) <= 4 and cod[:1].isdigit()):
                    cap, sub = cod, None
                    caps[cod] = Capitulo(cod, desc, None, 1, orden)
                else:
                    caps[cod] = Capitulo(cod, desc, cap, 2, orden)
                    sub = cod
            continue
        orden += 1
        cant = _num(g(f, "cant_origen")) or Decimal(0)
        precio = _num(g(f, "precio")) or Decimal(0)
        imp = _num(g(f, "imp_origen"))
        imp = (imp if imp is not None else cant * precio).quantize(Decimal("0.001"))
        an = _num(g(f, "imp_anterior"))
        ac = _num(g(f, "imp_actual"))
        # formato SIS: filas siguientes con marca AN / AC
        if det["modo"] == "sis":
            for f2 in filas[i + 1:i + 4]:
                mk = str(g(f2, "marca") or "").strip()
                if mk in ("AN", "AC"):
                    v = _num(g(f2, "imp_origen"))
                    v = v.quantize(Decimal("0.001")) if v is not None else None
                    if mk == "AN":
                        an = v
                    else:
                        ac = v
        if an is None and ac is not None:
            an = imp - ac
        if ac is None and an is not None:
            ac = imp - an
        pct = _num(g(f, "pct_origen")) or Decimal(0)
        c.lineas.append(Linea(codigo=cod, unidad=str(g(f, "unidad") or "") or None, titulo=_limpia(desc)[:200],
                              capitulo=cap or "?", subcapitulo=sub, pct_origen=pct, cant_origen=cant, precio=precio,
                              imp_origen=imp, cant_anterior=Decimal(0), imp_anterior=an if an is not None else Decimal(0),
                              cant_actual=Decimal(0), imp_actual=ac if ac is not None else Decimal(0), orden=orden))
        c._sin_anterior = an is None and ac is None
    for k in caps.values():
        ls = [l for l in c.lineas if (l.capitulo == k.codigo if k.nivel == 1 else l.subcapitulo == k.codigo)]
        k.tot_origen = sum((l.imp_origen for l in ls), Decimal(0))
        k.tot_anterior = sum((l.imp_anterior for l in ls), Decimal(0))
        k.tot_actual = sum((l.imp_actual for l in ls), Decimal(0))
    c.capitulos = sorted(caps.values(), key=lambda x: x.orden)
    c.totales = (sum((l.imp_origen for l in c.lineas), Decimal(0)), sum((l.imp_anterior for l in c.lineas), Decimal(0)),
                 sum((l.imp_actual for l in c.lineas), Decimal(0)))
    if getattr(c, "_sin_anterior", False) and c.lineas:
        c.avisos.append("La hoja solo trae el importe a origen: el «anterior» se toma de la certificación previa de la obra.")
    return c


def leer_excel(data: bytes, nombre: str) -> tuple[Certificacion, dict, list[list]]:
    import io as _io
    if nombre.lower().endswith(".csv"):
        import csv
        txt = data.decode("utf-8-sig", errors="replace")
        dial = csv.Sniffer().sniff(txt[:4000], delimiters=";,\t")
        filas = [list(r) for r in csv.reader(_io.StringIO(txt), dial)]
    else:
        from openpyxl import load_workbook
        wb = load_workbook(_io.BytesIO(data), read_only=True, data_only=True)
        hoja = next((n for n in wb.sheetnames if re.search(r"(?i)certif", n)), wb.sheetnames[0])
        filas = [list(r) for r in wb[hoja].iter_rows(values_only=True)]
    det = detectar_tabla(filas)
    return leer_tabla(filas, det), det, filas
