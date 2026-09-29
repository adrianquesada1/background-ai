"""
PREPARACIÓN DE LA CERTIFICACIÓN AL CLIENTE a partir de las certificaciones de proveedores.

Hasta ahora la certificación al cliente se hacía fuera (SIS/Presto) y se importaba hecha. Aquí se PROPONE:

- Estructura: la de la última certificación al cliente importada (capítulos, partidas, precios de venta, medición de
  presupuesto y lo certificado a origen).
- Avance: cada línea de contrato de subcontrata indica a qué partida del cliente corresponde (`cert_partida`). El % de
  avance de la partida es la media, ponderada por el importe contratado, del % a origen de sus líneas en las
  certificaciones de proveedor aprobadas hasta el periodo.
- Cantidad propuesta a origen = % de avance × medición de presupuesto de la partida (si una sola línea con la misma
  unidad la cubre, se toma directamente su cantidad a origen).
- Reglas de prudencia: nunca se propone menos de lo ya certificado al cliente (se avisa si la obra mide por debajo), ni
  más del 100 % del presupuesto sin marcarlo; las partidas sin subcontrata asociada mantienen lo certificado («sin
  medición»: las debe completar el jefe de obra).
- Resultado: venta del mes propuesta por partida y capítulo, frente al coste certificado a proveedores en el mismo mes
  (margen del mes antes de emitir). Se puede corregir a mano, guardar el borrador y descargar en Excel para SIS.
- Cuando se importe la certificación definitiva, se compara con la propuesta partida a partida.
"""
from __future__ import annotations

import io
import json
import re
from collections import defaultdict
from decimal import Decimal, ROUND_HALF_UP

from . import db

SCHEMA = """
CREATE TABLE IF NOT EXISTS cert_cliente_borradores (
    id INTEGER PRIMARY KEY, obra_id INTEGER NOT NULL, periodo TEXT NOT NULL, base_cert_id INTEGER, ajustes TEXT,
    estado TEXT DEFAULT 'borrador', creado_por TEXT, creado_en TEXT, actualizado_en TEXT, UNIQUE (obra_id, periodo));
"""


def init(con):
    con.executescript(SCHEMA)
    con.commit()


def _d(v) -> Decimal:
    try:
        return Decimal(str(v)) if v not in (None, "") else Decimal(0)
    except Exception:
        return Decimal(0)


def _q3(v: Decimal) -> Decimal:
    return v.quantize(Decimal("0.001"), rounding=ROUND_HALF_UP)


def _ud(u) -> str:
    u = re.sub(r"[^a-z0-9]", "", str(u or "").lower().replace("²", "2").replace("³", "3"))
    return {"ml": "m", "mts": "m", "u": "ud", "und": "ud", "uds": "ud", "pa": "pa"}.get(u, u)


def avance_lineas(con, obra_id: int, periodo: str) -> dict[str, list[dict]]:
    """Por partida del cliente: las líneas de contrato asociadas con su cantidad a origen certificada hasta el periodo."""
    out = defaultdict(list)
    for l in db.rows(con, """SELECT l.*, o.id AS oferta_id, COALESCE(pr.nombre, o.proveedor_nombre) AS proveedor FROM contrato_lineas l
                             JOIN ofertas o ON o.id=l.oferta_id LEFT JOIN proveedores pr ON pr.id=o.proveedor_id
                             WHERE o.obra_id=? AND o.estado='adjudicada' AND l.cert_partida IS NOT NULL AND l.cert_partida<>''""", (obra_id,)):
        c = db.one(con, """SELECT cl.cant_origen, cl.cant_mes FROM cert_prov_lineas cl JOIN certs_proveedor c ON c.id=cl.cert_id
                           WHERE cl.linea_id=? AND c.estado IN ('aprobada','facturada') AND c.periodo<=? ORDER BY c.numero DESC LIMIT 1""",
                   (l["id"], periodo))
        origen = _d(c["cant_origen"]) if c else Decimal(0)
        cant = _d(l["cantidad"])
        out[l["cert_partida"].strip()].append({"linea_id": l["id"], "proveedor": l["proveedor"], "descripcion": l["descripcion"],
                                                "unidad": l["unidad"], "cantidad": cant, "origen": origen,
                                                "pct": (origen / cant) if cant else Decimal(0), "peso": _d(l["importe_cents"])})
    return out


def proponer(con, obra_id: int, periodo: str, base_cert_id: int | None = None, ajustes: dict | None = None) -> dict:
    """Devuelve {'lineas': [...], 'capitulos': [...], 'totales': {...}, 'avisos': [...]} (importes en euros, Decimal)."""
    from . import obra_control, cert_proveedor
    base = db.one(con, "SELECT * FROM certificaciones WHERE id=?", (base_cert_id,)) if base_cert_id else obra_control.ultima_cert(con, obra_id)
    if not base:
        raise ValueError("La obra no tiene ninguna certificación al cliente importada: hace falta una para tomar la estructura y los precios.")
    ajustes = {str(k): v for k, v in (ajustes or {}).items()}
    av = avance_lineas(con, obra_id, periodo)
    lineas, avisos = [], []
    usados = set()
    for r in db.rows(con, "SELECT * FROM cert_lineas WHERE cert_id=? ORDER BY orden", (base["id"],)):
        precio = _d(r["precio"])
        ant = _d(r["cant_origen"])
        pres = _d(r["cant_presupuesto"])
        if not pres and _d(r["pct_origen"]):
            pres = _q3(ant / (_d(r["pct_origen"]) / 100))
        cod = (r["codigo"] or "").strip()
        clave = f"{r['id']}"
        fuente, aviso, prop = "", "", ant
        ls = av.get(cod, [])
        if ls:
            usados.add(cod)
            fuente = "; ".join(f"{x['proveedor']}: {x['origen']:.2f}/{x['cantidad']:.2f} {x['unidad'] or ''}" for x in ls)
            if len(ls) == 1 and _ud(ls[0]["unidad"]) == _ud(r["unidad"]) and _ud(r["unidad"]):
                prop = ls[0]["origen"]
            elif pres:
                peso = sum(x["peso"] for x in ls)
                pct = (sum(x["pct"] * x["peso"] for x in ls) / peso) if peso else sum(x["pct"] for x in ls) / len(ls)
                prop = _q3(pct * pres)
            else:
                aviso = "Sin medición de presupuesto para convertir el % de avance"
            if prop < ant:
                aviso = f"La obra mide {prop:.3f} a origen, menos de lo ya certificado ({ant:.3f}): se mantiene lo certificado"
                prop = ant
            elif pres and prop > pres:
                aviso = f"Supera la medición de presupuesto ({pres:.3f}): exceso de medición, revisar"
        elif precio:
            aviso = "Sin subcontrata asociada: se mantiene lo certificado (medir a mano si ha avanzado)"
        if clave in ajustes and ajustes[clave] not in (None, ""):
            prop = _d(ajustes[clave])
            fuente = (fuente + " · " if fuente else "") + "ajuste manual"
            if prop < ant:
                aviso = "Ajuste manual por debajo de lo certificado: descertifica"
        imp_origen = _q3(prop * precio)
        imp_ant = _d(r["origen_m"]) / 1000
        lineas.append({"id": r["id"], "capitulo": r["capitulo"], "subcapitulo": r["subcapitulo"], "codigo": cod, "titulo": r["titulo"],
                       "unidad": r["unidad"], "precio": precio, "cant_presupuesto": pres, "cant_anterior": ant, "cant_origen": prop,
                       "cant_mes": prop - ant, "pct_origen": (prop / pres * 100) if pres else None,
                       "importe_anterior": imp_ant, "importe_origen": imp_origen, "importe_mes": imp_origen - imp_ant,
                       "fuente": fuente, "aviso": aviso, "tipo": r["tipo"]})
    sin_partida = sorted(set(av) - usados)
    if sin_partida:
        avisos.append("Líneas de contrato asociadas a partidas que no existen en la certificación base: " + ", ".join(sin_partida[:15]))
    caps = defaultdict(lambda: {"importe_anterior": Decimal(0), "importe_origen": Decimal(0), "importe_mes": Decimal(0), "partidas": 0})
    for l in lineas:
        c = caps[l["capitulo"]]
        for k in ("importe_anterior", "importe_origen", "importe_mes"):
            c[k] += l[k]
        c["partidas"] += 1
    nombres = {r["codigo"]: r["nombre"] for r in db.rows(con, "SELECT codigo, nombre FROM cert_capitulos WHERE cert_id=?", (base["id"],))}
    capitulos = [{"capitulo": k, "nombre": nombres.get(k, ""), **v} for k, v in caps.items()]
    venta_mes = sum((l["importe_mes"] for l in lineas), Decimal(0))
    coste_mes = Decimal(cert_proveedor.coste_certificado_periodo(con, obra_id, periodo)) / 100
    return {"base": base, "periodo": periodo, "lineas": lineas, "capitulos": capitulos, "avisos": avisos,
            "totales": {"origen": sum((l["importe_origen"] for l in lineas), Decimal(0)),
                        "anterior": sum((l["importe_anterior"] for l in lineas), Decimal(0)), "mes": venta_mes,
                        "coste_proveedores_mes": coste_mes, "margen_mes": venta_mes - coste_mes,
                        "partidas_con_avance": sum(1 for l in lineas if l["cant_mes"] > 0),
                        "partidas_sin_medicion": sum(1 for l in lineas if l["aviso"].startswith("Sin subcontrata"))}}


def guardar_borrador(con, obra_id: int, periodo: str, base_cert_id: int, ajustes: dict, usuario: str) -> int:
    limpio = {str(k): str(v) for k, v in ajustes.items() if v not in (None, "")}
    with db.tx(con):
        con.execute("""INSERT INTO cert_cliente_borradores (obra_id, periodo, base_cert_id, ajustes, creado_por, creado_en, actualizado_en)
                       VALUES (?,?,?,?,?,?,?) ON CONFLICT(obra_id, periodo) DO UPDATE SET base_cert_id=excluded.base_cert_id,
                       ajustes=excluded.ajustes, actualizado_en=excluded.actualizado_en""",
                    (obra_id, periodo, base_cert_id, json.dumps(limpio), usuario, db.now_iso(), db.now_iso()))
        r = db.one(con, "SELECT id FROM cert_cliente_borradores WHERE obra_id=? AND periodo=?", (obra_id, periodo))
        db.audit(con, usuario, "cert_cliente_borrador", "obra", obra_id, {"periodo": periodo, "ajustes": len(limpio)})
    return r["id"]


def cargar_borrador(con, obra_id: int, periodo: str) -> dict | None:
    r = db.one(con, "SELECT * FROM cert_cliente_borradores WHERE obra_id=? AND periodo=?", (obra_id, periodo))
    if r:
        r["ajustes"] = json.loads(r["ajustes"] or "{}")
    return r


def excel(prop: dict) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    wb = Workbook()
    ws = wb.active
    ws.title = "Propuesta"
    cab = ["Capítulo", "Código", "Título", "Ud", "Precio", "Medición presupuesto", "Cant. anterior", "Cant. origen", "Cant. mes",
           "% origen", "Importe anterior", "Importe origen", "Importe mes", "Fuente (subcontratas)", "Aviso"]
    ws.append(cab)
    for c in ws[1]:
        c.font = Font(bold=True)
    amarillo = PatternFill("solid", fgColor="FFF2CC")
    for l in prop["lineas"]:
        ws.append([l["capitulo"], l["codigo"], l["titulo"], l["unidad"], float(l["precio"]), float(l["cant_presupuesto"]),
                   float(l["cant_anterior"]), float(l["cant_origen"]), float(l["cant_mes"]),
                   float(l["pct_origen"]) if l["pct_origen"] is not None else None, float(l["importe_anterior"]), float(l["importe_origen"]),
                   float(l["importe_mes"]), l["fuente"], l["aviso"]])
        if l["aviso"]:
            for c in ws[ws.max_row]:
                c.fill = amarillo
    for col, w in zip("ABCDEFGHIJKLMNO", [9, 12, 45, 6, 10, 12, 12, 12, 11, 9, 14, 14, 13, 40, 50]):
        ws.column_dimensions[col].width = w
    for fila in ws.iter_rows(min_row=2, min_col=11, max_col=13):
        for c in fila:
            c.number_format = '#,##0.00 "€"'
    ws2 = wb.create_sheet("Capítulos")
    ws2.append(["Capítulo", "Nombre", "Partidas", "Anterior", "Origen", "Mes"])
    for c in ws2[1]:
        c.font = Font(bold=True)
    for c in prop["capitulos"]:
        ws2.append([c["capitulo"], c["nombre"], c["partidas"], float(c["importe_anterior"]), float(c["importe_origen"]), float(c["importe_mes"])])
    t = prop["totales"]
    ws2.append([])
    ws2.append(["TOTAL", "", "", float(t["anterior"]), float(t["origen"]), float(t["mes"])])
    ws2.append(["Coste certificado a proveedores en el mes", "", "", "", "", float(t["coste_proveedores_mes"])])
    ws2.append(["Margen del mes (venta − coste subcontratas)", "", "", "", "", float(t["margen_mes"])])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def comparar_con_real(con, prop: dict, cert_id: int) -> list[dict]:
    """Propuesta frente a la certificación definitiva importada: diferencias por partida (código + capítulo)."""
    real = {(r["capitulo"], (r["codigo"] or "").strip()): r for r in db.rows(con, "SELECT * FROM cert_lineas WHERE cert_id=?", (cert_id,))}
    out = []
    for l in prop["lineas"]:
        r = real.get((l["capitulo"], l["codigo"]))
        if not r:
            continue
        dif = _d(r["cant_origen"]) - l["cant_origen"]
        if abs(dif) > Decimal("0.001"):
            out.append({"capitulo": l["capitulo"], "codigo": l["codigo"], "titulo": l["titulo"], "propuesta": l["cant_origen"],
                        "real": _d(r["cant_origen"]), "diferencia": dif, "diferencia_eur": _q3(dif * l["precio"])})
    return out
