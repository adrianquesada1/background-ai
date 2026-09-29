"""Exportación a Excel con importes como números exactos (Decimal) y formato contable."""
from __future__ import annotations

import io

import pandas as pd
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

from . import analytics, db
from .money import from_cents

EUR_FMT = '#,##0.00 "€";[Red]-#,##0.00 "€"'


def _eur(df: pd.DataFrame, cents_cols: dict[str, str]) -> pd.DataFrame:
    out = df.copy()
    for c, nuevo in cents_cols.items():
        if c in out:
            out[nuevo] = out[c].apply(lambda v: from_cents(int(v)) if pd.notna(v) else None)
            out = out.drop(columns=[c])
    return out


def exportar_excel(con, **f) -> bytes:
    docs = analytics.documentos_df(con, **f)
    lin = analytics.lineas_coste_df(con, **f)
    prov = analytics.por_proveedor(con, **f)
    ret = analytics.retenciones(con, **f)
    inc = pd.DataFrame(db.rows(con, """
        SELECT d.id AS documento, d.filename AS archivo, d.emisor_nombre AS proveedor, d.numero, i.severidad, i.codigo,
               i.mensaje, i.detalle, CASE i.resuelta WHEN 1 THEN 'Sí' ELSE 'No' END AS resuelta, i.resuelta_por, i.comentario
        FROM incidencias i JOIN documentos d ON d.id=i.documento_id ORDER BY d.id"""))
    hojas = {
        "Documentos": _eur(docs.drop(columns=["obra_id", "proveedor_id"], errors="ignore"),
                           {"base_c": "Base imponible", "iva_c": "IVA", "total_c": "Total factura", "irpf_c": "IRPF",
                            "ret_c": "Ret. garantía", "pagar_c": "A pagar"}),
        "Líneas por partida": _eur(lin.drop(columns=["obra_id", "proveedor_id", "partida_id"], errors="ignore"),
                                   {"importe_c": "Importe"}),
        "Proveedores": _eur(prov.drop(columns=["proveedor_id"], errors="ignore"),
                            {"base_c": "Base imponible", "ret_c": "Ret. garantía", "pagar_c": "A pagar", "pendiente_c": "Pendiente pago"}),
        "Retenciones garantía": _eur(ret, {"base_c": "Base imponible", "ret_c": "Retención"}) if not ret.empty else ret,
        "Incidencias": inc,
    }
    if f.get("obra_id"):
        pp = analytics.por_partida(con, f["obra_id"], **{k: v for k, v in f.items() if k != "obra_id"})
        hojas = {"Coste por partida": _eur(pp.drop(columns=["partida_id"], errors="ignore"),
                                           {"ppto_c": "Presupuesto coste", "coste_c": "Coste real",
                                            "desviacion_c": "Desviación", "extras_c": "Extras"}), **hojas}
    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as xw:
        for nombre, df in hojas.items():
            if df.empty:
                df = pd.DataFrame({"(sin datos)": []})
            df.to_excel(xw, sheet_name=nombre[:31], index=False)
            ws = xw.sheets[nombre[:31]]
            for cell in ws[1]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="2B2B2B")
                cell.alignment = Alignment(wrap_text=True, vertical="center")
            for i, col in enumerate(df.columns, start=1):
                letra = get_column_letter(i)
                ancho = min(60, max(10, int(df[col].map(lambda v: len(str(v)) if v is not None and v == v else 0).quantile(0.9)) + 2 if len(df) else 12, len(str(col)) + 2))
                ws.column_dimensions[letra].width = ancho
                if df[col].map(lambda v: hasattr(v, "quantize")).any():
                    for row in ws.iter_rows(min_row=2, min_col=i, max_col=i):
                        for c in row:
                            c.number_format = EUR_FMT
            ws.freeze_panes = "A2"
            ws.auto_filter.ref = ws.dimensions
    return buf.getvalue()
