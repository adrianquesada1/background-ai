"""
ESTUDIOS: medición → corte por industrial → petición de ofertas → comparativa homogénea → archivo de coste.

- Ingesta de mediciones en BC3 (FIEBDC-3: registros ~C conceptos, ~D descomposición, ~M mediciones, ~T textos) o Excel.
- Corte por industrial: el sistema propone el oficio de cada partida (mismas reglas que la rentabilidad por oficio);
  el técnico lo confirma o corrige.
- Pack de oferta por industrial: separata en Excel (código, ud, descripción, medición, precio a rellenar) y registro de a
  quién se pidió, cuándo y por qué medio.
- Las ofertas recibidas (la misma separata rellenada) se leen y se comparan partida a partida: huecos (partidas sin
  precio), condiciones y total, lado a lado. Referencia del histórico propio de Llorca (ofertas anteriores y compras).
- Archivo de coste con huecos: la mejor opción por partida y lo que falta por cubrir. El precio de venta lo fija Estudios.
"""
from __future__ import annotations

import io
import re
from datetime import date
from decimal import Decimal

from . import db
from .money import parse_amount, q2, to_cents

SCHEMA = """
CREATE TABLE IF NOT EXISTS estudios (
    id INTEGER PRIMARY KEY, nombre TEXT NOT NULL, cliente TEXT, obra_id INTEGER REFERENCES obras(id) ON DELETE SET NULL,
    estado TEXT DEFAULT 'abierto', origen TEXT, creado_por TEXT, creado_en TEXT, cerrado_por TEXT, cerrado_en TEXT, notas TEXT);
CREATE TABLE IF NOT EXISTS estudio_partidas (
    id INTEGER PRIMARY KEY, estudio_id INTEGER NOT NULL REFERENCES estudios(id) ON DELETE CASCADE,
    orden INTEGER, capitulo TEXT, codigo TEXT, unidad TEXT, descripcion TEXT, texto TEXT, medicion TEXT,
    precio_proyecto TEXT, oficio TEXT, oficio_origen TEXT DEFAULT 'propuesto');
CREATE INDEX IF NOT EXISTS ix_ep_est ON estudio_partidas(estudio_id);
CREATE TABLE IF NOT EXISTS estudio_solicitudes (
    id INTEGER PRIMARY KEY, estudio_id INTEGER NOT NULL REFERENCES estudios(id) ON DELETE CASCADE,
    oficio TEXT, empresa TEXT NOT NULL, email TEXT, medio TEXT, enviado_en TEXT, usuario TEXT,
    estado TEXT DEFAULT 'enviada', plazo TEXT, forma_pago TEXT, validez TEXT, exclusiones TEXT, notas TEXT, recibida_en TEXT);
CREATE TABLE IF NOT EXISTS estudio_precios (
    id INTEGER PRIMARY KEY, solicitud_id INTEGER NOT NULL REFERENCES estudio_solicitudes(id) ON DELETE CASCADE,
    partida_id INTEGER NOT NULL REFERENCES estudio_partidas(id) ON DELETE CASCADE, precio TEXT, UNIQUE (solicitud_id, partida_id));
"""


def init(con):
    con.executescript(SCHEMA)
    con.commit()


# ============================================================================ BC3 (FIEBDC-3)
def leer_bc3(data: bytes) -> list[dict]:
    """Partidas (conceptos con medición) con su capítulo, unidad, resumen, texto largo, medición y precio de proyecto."""
    try:
        txt = data.decode("cp850")
    except UnicodeDecodeError:
        txt = data.decode("latin-1", errors="replace")
    if "~V" not in txt[:2000] and "~C" not in txt:
        raise ValueError("No parece un archivo BC3 (FIEBDC-3).")
    conceptos, hijos, medic, textos, orden = {}, {}, {}, {}, []
    for reg in txt.split("~")[1:]:
        tipo, _, cuerpo = reg.partition("|")
        campos = cuerpo.rstrip("\r\n").split("|")
        if tipo == "C":
            cods = campos[0].split("\\")
            cod = cods[0].rstrip("#")
            conceptos[cod] = {"codigo": cod, "raiz": campos[0].endswith("##"), "capitulo_flag": campos[0].endswith("#"),
                              "unidad": campos[1] if len(campos) > 1 else "", "resumen": campos[2] if len(campos) > 2 else "",
                              "precio": (campos[3].split("\\")[0] if len(campos) > 3 else "") or "0"}
            orden.append(cod)
        elif tipo == "D":
            padre = campos[0].rstrip("#")
            trozos = (campos[1] if len(campos) > 1 else "").split("\\")
            hijos[padre] = [trozos[i].rstrip("#") for i in range(0, len(trozos) - 2, 3) if trozos[i]]
        elif tipo == "M":
            par = campos[0].split("\\")
            hijo = par[-1].rstrip("#") if len(par) > 1 else par[0].rstrip("#")
            total = campos[2] if len(campos) > 2 else ""
            try:
                medic[hijo] = Decimal(total.replace(",", ".")) if total.strip() else medic.get(hijo)
            except Exception:
                pass
        elif tipo == "T":
            textos[campos[0].rstrip("#")] = campos[1] if len(campos) > 1 else ""
    raiz = next((c for c in orden if conceptos[c]["raiz"]), orden[0] if orden else None)
    out, n = [], 0

    def recorrer(cod, cap, nivel):
        nonlocal n
        for h in hijos.get(cod, []):
            c = conceptos.get(h)
            if not c:
                continue
            if c["capitulo_flag"] or (h in hijos and h not in medic and nivel < 6 and not c["unidad"]):
                recorrer(h, f"{h} {c['resumen']}".strip(), nivel + 1)
            elif h in medic or c["unidad"]:
                n += 1
                out.append({"orden": n, "capitulo": cap, "codigo": h, "unidad": c["unidad"], "descripcion": c["resumen"],
                            "texto": textos.get(h, ""), "medicion": str(medic.get(h) or ""), "precio_proyecto": c["precio"]})
    if raiz:
        recorrer(raiz, conceptos[raiz]["resumen"], 0)
    return out


def leer_excel_mediciones(data: bytes, nombre: str) -> list[dict]:
    import pandas as pd
    df = pd.read_csv(io.BytesIO(data), sep=None, engine="python", dtype=str) if nombre.lower().endswith(".csv") \
        else pd.read_excel(io.BytesIO(data), dtype=str, header=None)
    filas = df.fillna("").values.tolist()
    cab_i, cols = None, {}
    for i, f in enumerate(filas[:30]):
        cab = [str(x).strip().lower() for x in f]
        m = {}
        for k, pat in (("codigo", r"^c[oó]d"), ("unidad", r"^(ud|uds|unidad|u)\.?$"), ("descripcion", r"descrip|resumen|concepto"),
                       ("medicion", r"medici|cantidad|cant\.?$"), ("precio", r"precio")):
            for j, h in enumerate(cab):
                if re.search(pat, h) and j not in m.values():
                    m[k] = j
                    break
        if "descripcion" in m and "medicion" in m:
            cab_i, cols = i, m
            break
    if cab_i is None:
        raise ValueError("No encuentro la cabecera: el Excel debe tener columnas de descripción y medición (y a ser posible código y unidad).")
    out, cap, n = [], "", 0
    for f in filas[cab_i + 1:]:
        g = lambda k: str(f[cols[k]]).strip() if k in cols and cols[k] < len(f) else ""  # noqa: E731
        desc, med = g("descripcion"), g("medicion")
        if not desc:
            continue
        if not med or not re.search(r"\d", med):
            cap = f"{g('codigo')} {desc}".strip()
            continue
        n += 1
        out.append({"orden": n, "capitulo": cap, "codigo": g("codigo") or f"P{n:04d}", "unidad": g("unidad"), "descripcion": desc[:300],
                    "texto": "", "medicion": str(parse_amount(med)), "precio_proyecto": str(parse_amount(g("precio"))) if g("precio") else ""})
    return out


def crear_estudio(con, nombre: str, cliente: str, partidas: list[dict], origen: str, usuario: str, obra_id=None) -> int:
    if not partidas:
        raise ValueError("No se han encontrado partidas con medición.")
    from .auditoria import OFICIOS, _puntua, norm
    kws = {n: {w for w in (norm(n).lower() + " " + norm(k).lower()).split() if len(w) >= 3} for n, t, k in OFICIOS if t == "DIRECTO"}
    with db.tx(con):
        cur = con.execute("INSERT INTO estudios (nombre, cliente, obra_id, origen, creado_por, creado_en) VALUES (?,?,?,?,?,?)",
                          (nombre, cliente, obra_id, origen, usuario, db.now_iso()))
        eid = cur.lastrowid
        for p in partidas:
            of, _ = _puntua(f"{p['descripcion']} {p['texto'][:300]} {p['capitulo']}", kws)
            con.execute("""INSERT INTO estudio_partidas (estudio_id, orden, capitulo, codigo, unidad, descripcion, texto, medicion,
                           precio_proyecto, oficio) VALUES (?,?,?,?,?,?,?,?,?,?)""",
                        (eid, p["orden"], p["capitulo"], p["codigo"], p["unidad"], p["descripcion"], p["texto"], p["medicion"],
                         p["precio_proyecto"], of))
        db.audit(con, usuario, "crear_estudio", "estudio", eid, {"partidas": len(partidas), "origen": origen})
    return eid


# ============================================================================ separatas y ofertas
def separata_excel(con, estudio_id: int, oficio: str, empresa: str = "") -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    e = db.one(con, "SELECT * FROM estudios WHERE id=?", (estudio_id,))
    ps = db.rows(con, "SELECT * FROM estudio_partidas WHERE estudio_id=? AND oficio=? ORDER BY orden", (estudio_id, oficio))
    wb = Workbook()
    ws = wb.active
    ws.title = "Oferta"
    ws["A1"] = f"Petición de oferta · {e['nombre']}"; ws["A1"].font = Font(bold=True, size=13)
    ws["A2"] = f"Industrial: {oficio}" + (f" · Empresa: {empresa}" if empresa else "")
    ws["A3"] = "Rellene SOLO la columna «Precio unitario» (€). No cambie los códigos. Indique plazo, forma de pago, validez y exclusiones abajo."
    cab = ["Código", "Ud", "Descripción", "Medición", "Precio unitario", "Importe"]
    for j, h in enumerate(cab, 1):
        c = ws.cell(row=5, column=j, value=h); c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="2E2E2E")
    for i, p in enumerate(ps, 6):
        ws.cell(row=i, column=1, value=p["codigo"]); ws.cell(row=i, column=2, value=p["unidad"])
        ws.cell(row=i, column=3, value=p["descripcion"])
        ws.cell(row=i, column=4, value=float(Decimal(p["medicion"] or 0)))
        ws.cell(row=i, column=5).fill = PatternFill("solid", fgColor="FFF7D6")
        ws.cell(row=i, column=6, value=f"=IF(E{i}=\"\",\"\",D{i}*E{i})")
    fin = 6 + len(ps)
    ws.cell(row=fin, column=5, value="TOTAL").font = Font(bold=True)
    ws.cell(row=fin, column=6, value=f"=SUM(F6:F{fin - 1})").font = Font(bold=True)
    for k, et in enumerate(["Plazo de ejecución", "Forma de pago", "Validez de la oferta", "Exclusiones / condiciones"], fin + 2):
        ws.cell(row=k, column=1, value=et).font = Font(bold=True)
    for col, w in zip("ABCDEF", [14, 6, 70, 12, 16, 16]):
        ws.column_dimensions[col].width = w
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def registrar_solicitud(con, estudio_id: int, oficio: str, empresa: str, email: str, medio: str, usuario: str) -> int:
    if not empresa.strip():
        raise ValueError("Indique la empresa.")
    ya = db.one(con, "SELECT id FROM estudio_solicitudes WHERE estudio_id=? AND oficio=? AND lower(empresa)=lower(?)",
                (estudio_id, oficio, empresa.strip()))
    if ya:
        raise ValueError("Ya se pidió oferta a esa empresa para ese industrial en este estudio.")
    with db.tx(con):
        cur = con.execute("INSERT INTO estudio_solicitudes (estudio_id, oficio, empresa, email, medio, enviado_en, usuario) VALUES (?,?,?,?,?,?,?)",
                          (estudio_id, oficio, empresa.strip(), email.strip(), medio, db.now_iso(), usuario))
        db.audit(con, usuario, "solicitud_oferta", "estudio", estudio_id, {"oficio": oficio, "empresa": empresa})
    return cur.lastrowid


def importar_oferta(con, solicitud_id: int, data: bytes, usuario: str) -> dict:
    """Lee la separata rellenada: precio por código y condiciones. Avisa de códigos desconocidos y partidas sin precio."""
    from openpyxl import load_workbook
    s = db.one(con, "SELECT * FROM estudio_solicitudes WHERE id=?", (solicitud_id,))
    ps = {p["codigo"]: p for p in db.rows(con, "SELECT * FROM estudio_partidas WHERE estudio_id=? AND oficio=?", (s["estudio_id"], s["oficio"]))}
    ws = load_workbook(io.BytesIO(data), data_only=True).active
    precios, desconocidos, cond = {}, [], {}
    for fila in ws.iter_rows(min_row=1, values_only=True):
        if not fila or fila[0] is None:
            continue
        cod = str(fila[0]).strip()
        if cod in ps:
            v = fila[4] if len(fila) > 4 else None
            if v not in (None, ""):
                try:
                    precios[cod] = Decimal(str(v)) if isinstance(v, (int, float)) else parse_amount(str(v))
                except Exception:
                    pass
        elif cod in ("Plazo de ejecución", "Forma de pago", "Validez de la oferta", "Exclusiones / condiciones"):
            cond[cod] = next((str(x) for x in fila[1:] if x not in (None, "")), "")
        elif re.match(r"^[A-Z0-9][\w.\-]{0,20}$", cod) and cod not in ("Código", "TOTAL"):
            desconocidos.append(cod)
    with db.tx(con):
        for cod, pr in precios.items():
            con.execute("INSERT INTO estudio_precios (solicitud_id, partida_id, precio) VALUES (?,?,?) "
                        "ON CONFLICT(solicitud_id, partida_id) DO UPDATE SET precio=excluded.precio", (solicitud_id, ps[cod]["id"], str(pr)))
        con.execute("UPDATE estudio_solicitudes SET estado='recibida', recibida_en=?, plazo=?, forma_pago=?, validez=?, exclusiones=? WHERE id=?",
                    (db.now_iso(), cond.get("Plazo de ejecución"), cond.get("Forma de pago"), cond.get("Validez de la oferta"),
                     cond.get("Exclusiones / condiciones"), solicitud_id))
        db.audit(con, usuario, "importar_oferta", "estudio", s["estudio_id"], {"empresa": s["empresa"], "precios": len(precios)})
    return {"precios": len(precios), "huecos": [c for c in ps if c not in precios], "desconocidos": desconocidos}


def referencia_historica(con, descripcion: str, unidad: str, excluir_estudio: int | None = None) -> Decimal | None:
    """Precio de referencia propio: mediana de ofertas anteriores de partidas parecidas (misma unidad) y de compras."""
    from rapidfuzz import fuzz
    cands = []
    for r in db.rows(con, """SELECT p.descripcion, p.unidad, pr.precio FROM estudio_precios pr JOIN estudio_partidas p ON p.id=pr.partida_id
                             WHERE p.estudio_id<>COALESCE(?, -1)""", (excluir_estudio,)):
        if (r["unidad"] or "").lower() == (unidad or "").lower() and fuzz.token_set_ratio(r["descripcion"], descripcion) >= 85:
            cands.append(Decimal(r["precio"]))
    if len(cands) < 2:
        for r in db.rows(con, "SELECT descripcion, unidad, precio_unitario FROM lineas WHERE precio_unitario IS NOT NULL LIMIT 20000"):
            if (r["unidad"] or "").lower() == (unidad or "").lower() and fuzz.token_set_ratio(r["descripcion"] or "", descripcion) >= 88:
                try:
                    cands.append(Decimal(r["precio_unitario"]))
                except Exception:
                    pass
    if not cands:
        return None
    cands.sort()
    return cands[len(cands) // 2]


def comparativa(con, estudio_id: int, oficio: str) -> dict:
    ps = db.rows(con, "SELECT * FROM estudio_partidas WHERE estudio_id=? AND oficio=? ORDER BY orden", (estudio_id, oficio))
    sols = db.rows(con, "SELECT * FROM estudio_solicitudes WHERE estudio_id=? AND oficio=? AND estado='recibida' ORDER BY empresa",
                   (estudio_id, oficio))
    precios = {(r["solicitud_id"], r["partida_id"]): Decimal(r["precio"]) for r in db.rows(con, """SELECT pr.* FROM estudio_precios pr
               JOIN estudio_solicitudes s ON s.id=pr.solicitud_id WHERE s.estudio_id=? AND s.oficio=?""", (estudio_id, oficio))}
    filas, totales, huecos = [], {s["id"]: Decimal(0) for s in sols}, {s["id"]: 0 for s in sols}
    for p in ps:
        med = Decimal(p["medicion"] or 0)
        f = {"codigo": p["codigo"], "ud": p["unidad"], "descripcion": p["descripcion"], "medicion": med,
             "referencia": referencia_historica(con, p["descripcion"], p["unidad"], estudio_id)}
        mejor = None
        for s in sols:
            pr = precios.get((s["id"], p["id"]))
            f[s["empresa"]] = pr
            if pr is None:
                huecos[s["id"]] += 1
            else:
                totales[s["id"]] += q2(pr * med)
                mejor = pr if mejor is None or pr < mejor else mejor
        f["mejor"] = mejor
        filas.append(f)
    return {"partidas": filas, "empresas": sols, "totales": totales, "huecos": huecos}


def archivo_coste(con, estudio_id: int) -> tuple[bytes, dict]:
    """Excel con el mejor precio por partida, la empresa, la referencia histórica y los huecos sin cubrir."""
    import pandas as pd
    e = db.one(con, "SELECT * FROM estudios WHERE id=?", (estudio_id,))
    oficios = [r["oficio"] for r in db.rows(con, "SELECT DISTINCT oficio FROM estudio_partidas WHERE estudio_id=? ORDER BY oficio", (estudio_id,))]
    filas, total, huecos = [], Decimal(0), 0
    for of in oficios:
        c = comparativa(con, estudio_id, of)
        for f in c["partidas"]:
            emp = None
            if f["mejor"] is not None:
                emp = next((s["empresa"] for s in c["empresas"] if f.get(s["empresa"]) == f["mejor"]), None)
            precio = f["mejor"] if f["mejor"] is not None else f["referencia"]
            origen = "Oferta" if f["mejor"] is not None else ("Histórico propio" if f["referencia"] is not None else "HUECO")
            imp = q2(precio * f["medicion"]) if precio is not None else None
            total += imp or 0
            huecos += origen == "HUECO"
            filas.append({"Oficio": of, "Código": f["codigo"], "Ud": f["ud"], "Descripción": f["descripcion"], "Medición": float(f["medicion"]),
                          "Precio coste": float(precio) if precio is not None else None, "Importe coste": float(imp) if imp is not None else None,
                          "Origen": origen, "Empresa": emp,
                          "Ref. histórica": float(f["referencia"]) if f["referencia"] is not None else None})
    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as xw:
        pd.DataFrame(filas).to_excel(xw, sheet_name="Archivo de coste", index=False)
    return buf.getvalue(), {"total": total, "huecos": huecos, "partidas": len(filas), "estudio": e["nombre"]}


def indicadores(con) -> dict:
    est = db.rows(con, "SELECT id, creado_en, cerrado_en FROM estudios")
    tiempos = [(date.fromisoformat(e["cerrado_en"][:10]) - date.fromisoformat(e["creado_en"][:10])).days for e in est if e["cerrado_en"]]
    compras = db.rows(con, "SELECT estudio_id, oficio, COUNT(*) n FROM estudio_solicitudes WHERE estado='recibida' GROUP BY estudio_id, oficio")
    tot_p = db.one(con, "SELECT COUNT(*) n FROM estudio_partidas")["n"]
    return {"estudios": len(est), "dias_medios": (sum(tiempos) / len(tiempos)) if tiempos else None,
            "ofertas_por_compra": (sum(c["n"] for c in compras) / len(compras)) if compras else None, "partidas": tot_p}
