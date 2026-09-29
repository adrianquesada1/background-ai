"""Pruebas de tesorería, pagos SEPA, cierre mensual y estudios (antes solo se comprobaban con scripts sueltos)."""
import io
import xml.etree.ElementTree as ET
from datetime import date, timedelta
from decimal import Decimal

import pytest

from core import db, tesoreria, pagos, cierres, estudios, ingesta


def _cert(con, obra, numero=1, fecha="2026-08-31", mes_m=100_000_000):
    cur = con.execute("INSERT INTO certificaciones (obra_id, numero, fecha, total_origen_m, total_anterior_m, total_actual_m) VALUES (?,?,?,?,?,?)",
                      (obra, numero, fecha, mes_m, 0, mes_m))
    con.commit()
    return cur.lastrowid


def _factura(con, obra, proveedor, num, pagar=100000, iban="ES7921000813610123456789", fecha="2026-08-10", venc=None):
    did, _ = ingesta.registrar_pdf(con, f"{num}.pdf", b"%PDF-1.4 " + num.encode() * 40, "t")
    con.execute("""UPDATE documentos SET tipo_documento='factura', estado='aprobada', proveedor_id=?, obra_id=?, numero=?, fecha=?,
                   fecha_vencimiento=?, base_imponible_cents=?, total_a_pagar_cents=?, iban=? WHERE id=?""",
                (proveedor, obra, num, fecha, venc or fecha, pagar, pagar, iban, did))
    con.commit()
    return did


# ============================================================================ tesorería
def test_factura_emitida_cobros_y_retencion(con, obra):
    cid = _cert(con, obra)
    fid = tesoreria.desde_certificacion(con, cid, "F1", "10", "5", "2026-09-01", "t")
    with pytest.raises(ValueError):
        tesoreria.desde_certificacion(con, cid, "F1", "10", "5", "2026-09-01", "t")         # una sola factura por certificación
    f = tesoreria.estado_factura(con, fid)
    assert f["base_cents"] == 10_000_000 and f["cuota_cents"] == 1_000_000 and f["retencion_cents"] == 500_000
    assert f["liquido_cents"] == 10_500_000
    with pytest.raises(ValueError):
        tesoreria.emitir(con, fid, "t", "gestor")                                            # falta NIF del cliente
    con.execute("UPDATE facturas_emitidas SET cliente_nif='A58818501' WHERE id=?", (fid,))
    with pytest.raises(ValueError):
        tesoreria.emitir(con, fid, "t", "consulta")
    assert tesoreria.emitir(con, fid, "t", "gestor") == "F1-2026/0001"
    fid2 = tesoreria.desde_certificacion(con, _cert(con, obra, 2, "2026-09-30"), "F1", "10", "5", "2026-10-01", "t")
    con.execute("UPDATE facturas_emitidas SET cliente_nif='A58818501' WHERE id=?", (fid2,))
    assert tesoreria.emitir(con, fid2, "t", "gestor") == "F1-2026/0002"                      # correlativa
    tesoreria.cobrar(con, fid, "2026-10-15", Decimal("50000"), "factura", "transferencia", "t")
    with pytest.raises(ValueError):
        tesoreria.cobrar(con, fid, "2026-10-16", Decimal("60000"), "factura", "transferencia", "t")   # más de lo pendiente
    assert tesoreria.estado_factura(con, fid)["pendiente_c"] == 5_500_000
    tesoreria.cobrar(con, fid, "2027-09-15", Decimal("5000"), "retencion", "transferencia", "t")
    assert db.one(con, "SELECT retencion_estado FROM facturas_emitidas WHERE id=?", (fid,))["retencion_estado"] == "cobrada"


def test_aging_radar_y_prevision(con, obra):
    hoy = date.today()
    fid = tesoreria.alta_historica(con, obra, "Cliente Viejo SA", "A58818501", "H-1", (hoy - timedelta(days=400)).isoformat(), "1000", "21", "5",
                                   "0", "t")
    ag = tesoreria.aging(con)
    assert ag[0]["tramo"] == "Más de 90 días"
    rad = tesoreria.radar_retenciones(con)
    assert rad[0]["situacion"].startswith("Vencida")
    assert "Devolución de retención" in tesoreria.borrador_reclamacion(con, fid, "Ana")
    prev = tesoreria.prevision_caja(con, 4, Decimal("100"))
    assert prev[0]["cobros"] == Decimal("1160.00") and prev[-1]["saldo"] == Decimal("1260.00")


# ============================================================================ pagos
def test_remesa_sepa_controles(con, obra, proveedor):
    db.set_setting(con, "sepa_nombre", "LLORCA GROUP HISPANIA SL")
    ok = _factura(con, obra, proveedor, "P-1")
    sin_iban = _factura(con, obra, proveedor, "P-2", iban=None)
    mal = _factura(con, obra, proveedor, "P-3", iban="ES0000000000000000000000")
    cands = {c["id"]: c for c in pagos.candidatas(con)}
    assert not cands[ok]["bloqueos"] and "sin IBAN" in cands[sin_iban]["bloqueos"] and "IBAN no válido" in cands[mal]["bloqueos"]
    with pytest.raises(ValueError, match="ordenante"):
        pagos.generar(con, [ok], "2026-09-30", "t", "gestor")
    db.set_setting(con, "sepa_iban", "ES9121000418450200051332")
    with pytest.raises(ValueError, match="No se pueden pagar"):
        pagos.generar(con, [ok, sin_iban], "2026-09-30", "t", "gestor")
    with pytest.raises(ValueError):
        pagos.generar(con, [ok], "2026-09-30", "t", "consulta")
    rid, xml = pagos.generar(con, [ok], "2026-09-30", "t", "gestor")
    ns = {"p": "urn:iso:std:iso:20022:tech:xsd:pain.001.001.03"}
    raiz = ET.fromstring(xml)
    assert raiz.find(".//p:GrpHdr/p:NbOfTxs", ns).text == "1" and raiz.find(".//p:GrpHdr/p:CtrlSum", ns).text == "1000.00"
    assert raiz.find(".//p:CdtTrfTxInf/p:CdtrAcct/p:Id/p:IBAN", ns).text == "ES7921000813610123456789"
    assert ok not in {c["id"] for c in pagos.candidatas(con)}                               # no puede ir en dos remesas
    pagos.confirmar(con, rid, "2026-09-30", "t")
    assert db.one(con, "SELECT pagada FROM documentos WHERE id=?", (ok,))["pagada"] == 1
    with pytest.raises(ValueError):
        pagos.confirmar(con, rid, "2026-09-30", "t")                                        # no se confirma dos veces
    # remesa anulada: la factura vuelve a estar disponible
    otra = _factura(con, obra, proveedor, "P-4")
    rid2, _ = pagos.generar(con, [otra], "2026-10-01", "t", "gestor")
    pagos.anular(con, rid2, "t")
    assert otra in {c["id"] for c in pagos.candidatas(con)}


def test_cambio_de_iban_sin_verificar_bloquea(con, obra, proveedor):
    did = _factura(con, obra, proveedor, "P-9")
    con.execute("INSERT INTO incidencias (documento_id, codigo, severidad, mensaje) VALUES (?, 'C11', 'critica', 'Cambio de cuenta')", (did,))
    con.commit()
    c = next(c for c in pagos.candidatas(con) if c["id"] == did)
    assert "cambio de cuenta bancaria sin verificar" in c["bloqueos"]


# ============================================================================ cierre mensual
def test_cierre_bloquea_y_reapertura(con, obra, proveedor):
    _cert(con, obra, 1, "2026-08-31")
    did = _factura(con, obra, proveedor, "C-1", fecha="2026-08-10")
    lc = {x["punto"]: x for x in cierres.lista_comprobacion(con, obra, "2026-08")}
    assert lc["Certificación del mes importada"]["ok"] and lc["Facturas del mes aprobadas o rechazadas"]["ok"]
    with pytest.raises(ValueError):
        cierres.cerrar(con, obra, "2026-08", "t", "jefe_obra")
    cierres.cerrar(con, obra, "2026-08", "t", "gestor")
    assert cierres.cerrado(con, obra, "2026-08-15")
    with pytest.raises(cierres.PeriodoCerrado):
        cierres.comprobar_documento(con, did)
    with pytest.raises(ValueError):
        cierres.reabrir(con, obra, "2026-08", "t", "gestor", "error")                        # solo Dirección o administrador
    with pytest.raises(ValueError):
        cierres.reabrir(con, obra, "2026-08", "t", "direccion", " ")
    cierres.reabrir(con, obra, "2026-08", "d", "direccion", "Factura tardía de agosto")
    assert not cierres.cerrado(con, obra, "2026-08-15")


def test_cierre_con_pendientes_exige_motivo(con, obra, proveedor):
    did = _factura(con, obra, proveedor, "C-2", fecha="2026-07-10")
    con.execute("UPDATE documentos SET estado='pendiente_revision' WHERE id=?", (did,))
    con.commit()
    with pytest.raises(ValueError, match="pendientes"):
        cierres.cerrar(con, obra, "2026-07", "t", "gestor")
    cierres.cerrar(con, obra, "2026-07", "t", "gestor", comentario="Se cierra con una factura en disputa", forzar=True)
    assert cierres.cerrado(con, obra, "2026-07-01")


# ============================================================================ estudios
BC3 = ("~V|SOFT S.A.|FIEBDC-3/2016|Presto|\r\n"
       "~C|OBRA##||Edificio prueba|0|\r\n~C|01#||ALBAÑILERÍA|0|\r\n~C|01.01|m2|Fábrica de ladrillo cerámico|25.5|\r\n"
       "~C|01.02|m2|Tabique de placa de yeso laminado|18|\r\n~C|02#||PINTURA|0|\r\n~C|02.01|m2|Pintura plástica lisa en paramentos|4.2|\r\n"
       "~D|OBRA##|01#\\1\\1\\02#\\1\\1\\|\r\n~D|01#|01.01\\1\\1\\01.02\\1\\1\\|\r\n~D|02#|02.01\\1\\1\\|\r\n"
       "~M|01#\\01.01|1|120.50|\r\n~M|01#\\01.02|1|300|\r\n~M|02#\\02.01|1|900|\r\n~T|01.01|Fábrica de ladrillo hueco doble|\r\n").encode("cp850")


def test_estudio_bc3_separata_ofertas_y_coste(con):
    ps = estudios.leer_bc3(BC3)
    assert [p["codigo"] for p in ps] == ["01.01", "01.02", "02.01"] and ps[0]["medicion"] == "120.50"
    with pytest.raises(ValueError):
        estudios.leer_bc3(b"esto no es un bc3")
    eid = estudios.crear_estudio(con, "Edificio prueba", "Promotora", ps, "bc3", "t")
    oficio = db.one(con, "SELECT oficio FROM estudio_partidas WHERE codigo='01.01' AND estudio_id=?", (eid,))["oficio"]
    con.execute("UPDATE estudio_partidas SET oficio=? WHERE estudio_id=? AND codigo IN ('01.01','01.02')", (oficio, eid))
    con.commit()
    sep = estudios.separata_excel(con, eid, oficio, "Albañiles A")
    s1 = estudios.registrar_solicitud(con, eid, oficio, "Albañiles A", "a@a.es", "correo", "t")
    s2 = estudios.registrar_solicitud(con, eid, oficio, "Albañiles B", "b@b.es", "correo", "t")
    with pytest.raises(ValueError):
        estudios.registrar_solicitud(con, eid, oficio, "albañiles a", "a@a.es", "correo", "t")   # sin duplicados
    from openpyxl import load_workbook

    def rellenar(precios):
        wb = load_workbook(io.BytesIO(sep))
        ws = wb.active
        for fila in ws.iter_rows(min_row=6):
            if fila[0].value in precios:
                fila[4].value = precios[fila[0].value]
        buf = io.BytesIO()
        wb.save(buf)
        return buf.getvalue()
    r = estudios.importar_oferta(con, s1, rellenar({"01.01": 20, "01.02": 15}), "t")
    assert r["precios"] == 2 and not r["huecos"]
    r = estudios.importar_oferta(con, s2, rellenar({"01.01": 19}), "t")
    assert r["huecos"] == ["01.02"]
    comp = estudios.comparativa(con, eid, oficio)
    assert comp["partidas"][0]["mejor"] == Decimal("19")
    xls, res = estudios.archivo_coste(con, eid)
    assert xls[:2] == b"PK" and res["partidas"] == 3
